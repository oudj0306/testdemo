"""文档类型验证、页面归一化与 OCR 输入准备。"""

from __future__ import annotations

import hashlib
import io
import tempfile
import uuid
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from functools import wraps
from pathlib import Path
from threading import RLock
from typing import Any

from PIL import Image, ImageOps

from grading_agent.config import Settings
from grading_agent.services.documents.legacy_word import extract_doc_text

ALLOWED_MEDIA = {
    "application/msword": ".doc",
    "application/pdf": ".pdf",
    "image/jpeg": ".jpg",
    "image/png": ".png",
    "image/webp": ".webp",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
}
DOCX_IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}

_document_lock = RLock()


def serialized_document_io(function):
    # PDFium access must not overlap across background workers.
    @wraps(function)
    def wrapped(*args, **kwargs):
        with _document_lock:
            return function(*args, **kwargs)

    return wrapped


@dataclass
class NormalizedPage:
    page_no: int
    path: Path
    width: int
    height: int
    quality_state: str
    quality_report: dict[str, Any]


@dataclass
class IngestedDocument:
    source_hash: str
    source_path: Path
    media_type: str
    pages: list[NormalizedPage]
    source_page_count: int | None = None


def sniff_media_type(content: bytes, declared: str | None) -> str:
    if content.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        import olefile

        try:
            with olefile.OleFileIO(io.BytesIO(content)) as compound:
                if compound.exists("WordDocument"):
                    return "application/msword"
        except OSError:
            pass
        raise ValueError("不是有效的 Word DOC 文件")
    if content.startswith(b"%PDF-"):
        return "application/pdf"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image/webp"
    if content.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        except zipfile.BadZipFile:
            pass
    raise ValueError("支持 DOC、DOCX、PDF、JPEG、PNG、WebP；文件内容与支持格式不符")


def _quality(width: int, height: int) -> tuple[str, dict[str, Any]]:
    shortest = min(width, height)
    if shortest < 900:
        return "low_resolution", {
            "short_edge": shortest,
            "recommended_action": "resubmit_higher_resolution",
        }
    return "ready", {"short_edge": shortest}


def _save_image_page(image: Image.Image, destination: Path, page_no: int) -> NormalizedPage:
    normalized = ImageOps.exif_transpose(image).convert("RGB")
    normalized.save(destination, format="PNG", optimize=True)
    state, report = _quality(*normalized.size)
    return NormalizedPage(page_no, destination, normalized.width, normalized.height, state, report)


@serialized_document_io
def ingest_document(
    content: bytes,
    declared_media_type: str | None,
    settings: Settings,
    render_pages: bool = True,
    render_word_pages: bool = False,
) -> IngestedDocument:
    media_type = sniff_media_type(content, declared_media_type)
    source_hash = hashlib.sha256(content).hexdigest()
    source_path = settings.upload_dir / f"{source_hash}{ALLOWED_MEDIA[media_type]}"
    if not source_path.exists():
        source_path.write_bytes(content)

    document_id = uuid.uuid4().hex
    pages: list[NormalizedPage] = []
    page_count = None
    if media_type == "application/pdf":
        try:
            import pypdfium2 as pdfium
        except ImportError as exc:
            raise RuntimeError("PDF 支持需要安装 pypdfium2") from exc
        with pdfium.PdfDocument(str(source_path)) as pdf:
            page_count = len(pdf)
            if page_count == 0 or page_count > 100:
                raise ValueError("PDF 页数必须在 1 到 100 之间")
            if render_pages:
                for index in range(page_count):
                    page = pdf[index]
                    bitmap = page.render(scale=2.0)
                    try:
                        image = bitmap.to_pil()
                        destination = settings.page_dir / f"{document_id}-{index + 1}.png"
                        pages.append(_save_image_page(image, destination, index + 1))
                    finally:
                        bitmap.close()
                        page.close()
    elif media_type.startswith("image/"):
        try:
            with Image.open(source_path) as image:
                destination = settings.page_dir / f"{document_id}-1.png"
                pages.append(_save_image_page(image, destination, 1))
        except Exception as exc:
            raise ValueError("图片无法解码") from exc
    elif (
        media_type
        in {
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        }
        and render_pages
        and render_word_pages
    ):
        from grading_agent.services.documents.document_renderer import word_to_pdf

        try:
            pdf_path = word_to_pdf(source_path, settings.data_dir)
            rendered = ingest_document(
                pdf_path.read_bytes(), "application/pdf", settings, render_pages=True
            )
            pages, page_count = rendered.pages, rendered.source_page_count
        except ValueError:
            # Word text extraction remains usable without an installed renderer.
            pages = []
    else:
        # DOCX is retained as-is; text and embedded media are extracted separately.
        pages = []
    return IngestedDocument(source_hash, source_path, media_type, pages, page_count)


def crop_question_regions(
    pages: list[NormalizedPage],
    questions: list[dict[str, Any]],
    settings: Settings,
) -> dict[str, Path]:
    page_map = {page.page_no: page for page in pages}
    crops: dict[str, Path] = {}
    for question in questions:
        region = question.get("rule_config", {}).get("region")
        if not region:
            continue
        from grading_agent.schemas import RegionInput

        try:
            region = RegionInput.model_validate(region).model_dump(exclude_none=True)
        except ValueError:
            continue
        page = page_map.get(int(region.get("page", 1)))
        if not page:
            continue
        with Image.open(page.path) as image:
            if (
                region.get("page_aspect")
                and abs(image.width / image.height / region["page_aspect"] - 1) > 0.1
            ):
                continue  # likely wrong page/camera geometry: use visual fallback
            x = int(float(region["x"]) * image.width)
            y = int(float(region["y"]) * image.height)
            width = int(float(region["width"]) * image.width)
            height = int(float(region["height"]) * image.height)
            if width <= 0 or height <= 0:
                continue
            crop = image.crop((x, y, min(x + width, image.width), min(y + height, image.height)))
            path = settings.crop_dir / f"{uuid.uuid4().hex}-{question['id']}.png"
            crop.save(path, format="PNG", optimize=True)
            crops[question["id"]] = path
    return crops


@serialized_document_io
def pdf_visual_pages(
    path: Path, existing_pages: list[NormalizedPage] | None = None
) -> list[tuple[int, bytes]]:
    """Return every PDF page as an image for vision APIs without PDF support.

    Reuse pages rendered during ingestion when available. Otherwise render in
    memory, so a retry does not create another set of persisted page files.
    """
    if existing_pages:
        return [(page.page_no, page.path.read_bytes()) for page in existing_pages]
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        raise RuntimeError("PDF 支持需要安装 pypdfium2") from exc
    images: list[tuple[int, bytes]] = []
    with pdfium.PdfDocument(str(path)) as pdf:
        if not 1 <= len(pdf) <= 100:
            raise ValueError("PDF 页数必须在 1 到 100 之间")
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=2.0)
            try:
                with io.BytesIO() as stream:
                    bitmap.to_pil().convert("RGB").save(stream, format="PNG", optimize=True)
                    images.append((index + 1, stream.getvalue()))
            finally:
                bitmap.close()
                page.close()
    return images


class OCRAdapter:
    def __init__(self, provider: str, device: str | None = None):
        self.provider = provider
        from grading_agent.config import get_settings

        self.device = device or get_settings().ocr_device

    def extract(self, path: Path) -> tuple[str, float | None, str]:
        if self.provider != "paddleocr":
            return "", None, "manual"
        from grading_agent.services.recognition.ocr.runtime import recognize

        return recognize(path, self.device)


def docx_embedded_images(path: Path) -> list[tuple[str, bytes]]:
    images: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            suffix = Path(name).suffix.lower()
            if name.startswith("word/media/") and suffix in DOCX_IMAGE_TYPES:
                images.append((DOCX_IMAGE_TYPES[suffix], archive.read(name)))
    return images


@serialized_document_io
def extract_reference_text(document: IngestedDocument, ocr: OCRAdapter) -> tuple[str, str]:
    """Extract reusable answer-key/rubric text once; prefer native PDF text over OCR."""
    if document.media_type == "application/msword":
        return extract_doc_text(document.source_path), "doc_text_partial_layout"
    if document.media_type == "application/pdf":
        try:
            import pypdfium2 as pdfium

            pdf = pdfium.PdfDocument(str(document.source_path))
            native_pages: list[str] = []
            for index in range(len(pdf)):
                text_page = pdf[index].get_textpage()
                native_pages.append(text_page.get_text_range())
            native_text = "\n\n".join(value.strip() for value in native_pages if value.strip())
            if native_text:
                return native_text, "native_pdf_text"
        except Exception:
            pass
    if (
        document.media_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        from docx import Document

        word = Document(str(document.source_path))
        blocks = [paragraph.text.strip() for paragraph in word.paragraphs if paragraph.text.strip()]
        for table in word.tables:
            for row in table.rows:
                line = "\t".join(cell.text.strip() for cell in row.cells)
                if line.strip():
                    blocks.append(line)
        try:
            with zipfile.ZipFile(document.source_path) as archive:
                root = ET.fromstring(archive.read("word/document.xml"))
            math_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
            for equation in root.iter(f"{math_namespace}oMath"):
                formula = "".join(
                    node.text or "" for node in equation.iter() if node.tag.endswith("}t")
                )
                if formula.strip():
                    blocks.append(f"[公式]{formula.strip()}")
        except (KeyError, ET.ParseError, zipfile.BadZipFile):
            pass
        image_texts: list[str] = []
        for content_type, blob in docx_embedded_images(document.source_path):
            suffix = ".png" if content_type == "image/png" else ".jpg"
            temp_path: Path | None = None
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as handle:
                    handle.write(blob)
                    temp_path = Path(handle.name)
                value, _, source = ocr.extract(temp_path)
                if value.strip():
                    image_texts.append(value.strip())
            finally:
                if temp_path and temp_path.exists():
                    temp_path.unlink()
        combined = "\n".join(blocks + image_texts)
        if combined:
            state = "docx_text+embedded_image_ocr" if image_texts else "docx_text"
            return combined, state
        return "", "stored_only"
    extracted: list[str] = []
    sources: set[str] = set()
    for page in document.pages:
        text, _, source = ocr.extract(page.path)
        sources.add(source)
        if text.strip():
            extracted.append(text.strip())
    if extracted:
        return "\n\n".join(extracted), "+".join(sorted(sources))
    return "", "stored_only"

"""Read-only preview, audit and reproducible reports; no model tokens required."""

import json
from html import escape
from pathlib import Path
from statistics import mean

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse, HTMLResponse

from grading_agent.config import get_settings
from grading_agent.db import transaction, utc_now
from grading_agent.repository import get_assessment, get_submission, rows_dict

router = APIRouter()


@router.get("/api/preview/{kind}/{record_id}")
def preview(kind: str, record_id: str):
    sources = {
        "submission": ("submissions", "source_path"),
        "document": ("assessment_documents", "source_path"),
        "page": ("submission_pages", "image_path"),
    }
    if kind not in sources:
        raise HTTPException(404, "文件类型不存在")
    table, column = sources[kind]
    with transaction() as db:
        row = db.execute(f"SELECT * FROM {table} WHERE id=?", (record_id,)).fetchone()
    if not row:
        raise HTTPException(404, "文件不存在")
    path = Path(row[column]).resolve()
    if not path.is_relative_to(get_settings().data_dir) or not path.is_file():
        raise HTTPException(404, "文件不存在或不在数据目录内")
    if path.suffix.lower() == ".doc":
        from grading_agent.services.documents.legacy_word import extract_doc_text

        return HTMLResponse(
            '<meta charset="utf-8"><style>body{font:16px sans-serif;padding:24px}pre{white-space:pre-wrap;line-height:1.8}</style><p>DOC 内容预览（包含文本框，不保留原始排版和绘图）</p><pre>'
            + escape(extract_doc_text(path))
            + "</pre>",
            headers={"Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'"},
        )
    if path.suffix.lower() == ".docx":
        from docx import Document

        word = Document(path)
        blocks = "".join(f"<p>{escape(p.text)}</p>" for p in word.paragraphs)
        blocks += "".join(
            "<table>"
            + "".join(
                "<tr>" + "".join(f"<td>{escape(c.text)}</td>" for c in r.cells) + "</tr>"
                for r in t.rows
            )
            + "</table>"
            for t in word.tables
        )
        return HTMLResponse(
            '<meta charset="utf-8"><style>body{font:16px sans-serif;padding:24px}td{border:1px solid #ddd;padding:8px}table{border-collapse:collapse}</style><p>Word 内容预览（非原始排版；公式及浮动图片请查看原件）</p>'
            + blocks,
            headers={"Content-Security-Policy": "default-src 'none'; style-src 'unsafe-inline'"},
        )
    return FileResponse(
        path,
        media_type="application/pdf"
        if path.suffix == ".pdf"
        else "image/png"
        if kind == "page"
        else row["media_type"],
        headers={"X-Content-Type-Options": "nosniff"},
    )


@router.get("/api/submissions/{submission_id}/history")
def history(submission_id: str):
    if not get_submission(submission_id):
        raise HTTPException(404, "答卷不存在")
    with transaction() as db:
        grades = rows_dict(
            db.execute(
                "SELECT * FROM grade_items WHERE submission_id=? ORDER BY created_at DESC",
                (submission_id,),
            ).fetchall()
        )
        for g in grades:
            g["reviews"] = rows_dict(
                db.execute(
                    "SELECT * FROM review_events WHERE grade_item_id=? ORDER BY created_at,rowid",
                    (g["id"],),
                ).fetchall()
            )
    return grades


@router.get("/api/assessments/{assessment_id}/report")
def report(assessment_id: str):
    assessment = get_assessment(assessment_id)
    if not assessment:
        raise HTTPException(404, "试卷不存在")
    with transaction() as db:
        ids = [
            r[0]
            for r in db.execute(
                "SELECT id FROM submissions WHERE assessment_id=? ORDER BY created_at",
                (assessment_id,),
            )
        ]
    submissions = [get_submission(i) for i in ids]
    students = []
    for s in submissions:
        summary = s["score_summary"]
        final = bool(s["grade_items"]) and summary["unresolved_count"] == 0
        students.append(
            {
                "id": s["id"],
                "student": s["student_ref"],
                "confirmed_score": summary["current_score"],
                "final_score": summary["current_score"] if final else None,
                "pending": summary["unresolved_count"],
                "review_recommended": summary["review_recommended_count"],
            }
        )
    totals = [s["final_score"] for s in students if s["final_score"] is not None]
    questions = []
    for q in assessment["questions"]:
        values = []
        for s in submissions:
            for g in s["grade_items"]:
                if g["question_id"] == q["id"] and g["effective_score"] is not None:
                    values.append(g["effective_score"])
        questions.append(
            {
                "question_no": q["question_no"],
                "max_score": q["max_score"],
                "sample_count": len(values),
                "average": round(mean(values), 2) if values else None,
                "score_rate": round(mean(values) / q["max_score"] * 100, 1) if values else None,
            }
        )
    auto_count = sum(
        g["grader_kind"] == "llm" and g["final_score"] is None and g["effective_score"] is not None
        for s in submissions
        for g in s["grade_items"]
    )
    rule_count = sum(
        g["grader_kind"] != "llm" and g["final_score"] is None and g["effective_score"] is not None
        for s in submissions
        for g in s["grade_items"]
    )
    human_count = sum(g["final_score"] is not None for s in submissions for g in s["grade_items"])
    pending_count = sum(s["pending"] for s in students)
    review_recommended_count = sum(s["review_recommended"] for s in students)
    total_max = sum(float(q["max_score"]) for q in assessment["questions"])
    distribution = [{"label": f"{i * 20}-{(i + 1) * 20}%", "count": 0} for i in range(5)]
    if total_max > 0:
        for s in students:
            if s["final_score"] is None:
                continue
            pct = float(s["final_score"]) / total_max * 100
            distribution[min(4, int(pct // 20))]["count"] += 1
    return {
        "title": assessment["title"],
        "generated_at": utc_now(),
        "submission_count": len(students),
        "finalized_count": len(totals),
        "average": round(mean(totals), 2) if totals else None,
        "highest": max(totals) if totals else None,
        "lowest": min(totals) if totals else None,
        "total_max": round(total_max, 2),
        "students": students,
        "questions": questions,
        "score_distribution": distribution,
        "auto_stable_question_count": auto_count,
        "rule_passed_question_count": rule_count,
        "human_reviewed_question_count": human_count,
        "pending_question_count": pending_count,
        "review_recommended_question_count": review_recommended_count,
        "method": f"每份答卷计一个样本；模型或规则给出数值后即形成当前成绩，人工复核用于可选纠错，最新复核分覆盖机器分。缺少分数的题不按零分计入。本卷模型评分 {auto_count} 题，规则评分 {rule_count} 题，人工复核 {human_count} 题，建议复核 {review_recommended_count} 题，仍缺分 {pending_count} 题。",
    }


@router.get("/api/assessments/{assessment_id}/report.html")
def report_html(assessment_id: str):
    data = report(assessment_id)
    students = "".join(
        f"<tr><td>{escape(s['student'])}</td><td>{s['final_score'] if s['final_score'] is not None else '缺少分数'}</td><td>{s['review_recommended']}</td><td>{s['pending']}</td></tr>"
        for s in data["students"]
    )
    chart_data = json.dumps(
        {
            "dist_labels": [d["label"] for d in data["score_distribution"]],
            "dist_counts": [d["count"] for d in data["score_distribution"]],
            "q_labels": [f"第{q['question_no']}题" for q in data["questions"]],
            "q_rates": [q["score_rate"] for q in data["questions"]],
            "mode": [
                data["auto_stable_question_count"],
                data["rule_passed_question_count"],
                data["human_reviewed_question_count"],
                data["pending_question_count"],
            ],
        },
        ensure_ascii=False,
    ).replace("<", "\\u003c")
    try:
        chart_js = (Path(__file__).resolve().parent / "static" / "chart.umd.min.js").read_text(
            encoding="utf-8"
        )
    except OSError:
        chart_js = ""
    css = (
        "body{font:16px sans-serif;max-width:960px;margin:40px auto;padding:20px;color:#233}"
        "table{border-collapse:collapse;width:100%;margin:20px 0}td,th{border:1px solid #ddd;padding:12px;text-align:left}p{line-height:1.8}"
        ".charts{display:grid;grid-template-columns:1fr 1fr;gap:18px;margin:16px 0}"
        ".chart-box{border:1px solid #e1e7e2;border-radius:10px;padding:16px;background:#fff;margin:12px 0}"
        ".chart-box h3{margin:0 0 12px}.chart-box canvas{max-height:260px}"
    )
    script = (
        f"<script>{chart_js}</script><script>const D={chart_data};"
        "new Chart(document.getElementById('distChart'),{type:'bar',data:{labels:D.dist_labels,datasets:[{data:D.dist_counts,backgroundColor:'#4e8a7a'}]},options:{plugins:{legend:{display:false}},scales:{y:{beginAtZero:true,ticks:{precision:0}}}}});"
        "new Chart(document.getElementById('modeChart'),{type:'doughnut',data:{labels:['模型评分','规则评分','人工复核','缺少分数'],datasets:[{data:D.mode,backgroundColor:['#5275d9','#78a999','#c7a575','#dfe5ef']}]}});"
        "new Chart(document.getElementById('qRateChart'),{type:'line',data:{labels:D.q_labels,datasets:[{label:'得分率 %',data:D.q_rates,borderColor:'#4e8a7a',backgroundColor:'#4e8a7a',fill:false,tension:0.3}]},options:{scales:{y:{beginAtZero:true,max:100}}}});"
        "</script>"
    )
    body = (
        f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>批改分析报告</title><style>{css}</style>'
        f"<h1>{escape(data['title'])} · 分析报告</h1>"
        f"<p>生成时间：{escape(data['generated_at'])}</p>"
        f"<p>答卷 {data['submission_count']} 份，已定分 {data['finalized_count']} 份；平均分 {data['average'] if data['average'] is not None else '暂无'}（满分 {data['total_max']}）。</p>"
        f"<p>{escape(data['method'])}</p>"
        f'<div class="charts"><div class="chart-box"><h3>成绩分布</h3><canvas id="distChart"></canvas></div><div class="chart-box"><h3>批改方式</h3><canvas id="modeChart"></canvas></div></div>'
        f'<div class="chart-box"><h3>逐题得分率趋势</h3><canvas id="qRateChart"></canvas></div>'
        f"<h2>学生成绩</h2><table><tr><th>学生</th><th>当前成绩</th><th>建议复核</th><th>缺少分数</th></tr>{students}</table>"
        + script
        + "</html>"
    )
    return HTMLResponse(
        body, headers={"Content-Disposition": 'attachment; filename="grading-report.html"'}
    )
